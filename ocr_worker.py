#!/usr/bin/env python3
"""Linux-capable OCR worker used by server.mjs.

Input: a JSON array of source file descriptors on stdin.
Output: one JSON object per recognized page on stdout, matching ocr.swift.
"""

import json
import os
import subprocess
import sys
from pathlib import Path


def emit(value):
    print(json.dumps(value, ensure_ascii=False), flush=True)


def average(values):
    return sum(values) / len(values) if values else 0.0


def read_text_file(path):
    data = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "gb18030", "latin-1"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            pass
    return ""


def extract_docx(path):
    from docx import Document

    document = Document(path)
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                lines.append("\t".join(cells))
    return "\n".join(lines)


def text_from_ocr_result(result):
    payload = getattr(result, "json", {})
    if callable(payload):
        payload = payload()
    if isinstance(payload, str):
        try:
            payload = json.loads(payload)
        except json.JSONDecodeError:
            payload = {}
    if not isinstance(payload, dict):
        return "", 0.0
    data = payload.get("res", payload)
    if not isinstance(data, dict):
        return "", 0.0
    texts = data.get("rec_texts")
    scores = data.get("rec_scores")
    boxes = data.get("rec_boxes")
    texts = list(texts) if texts is not None else []
    scores = list(scores) if scores is not None else []
    boxes = list(boxes) if boxes is not None else []
    rows = []
    for index, text in enumerate(texts):
        text = str(text or "").strip()
        if not text:
            continue
        score = float(scores[index]) if index < len(scores) else 0.0
        box = boxes[index] if index < len(boxes) else []
        try:
            # PaddleOCR boxes are [left, top, right, bottom]. Sorting provides
            # stable reading order even when a model returns boxes unordered.
            top = float(box[1])
            left = float(box[0])
        except (TypeError, ValueError, IndexError):
            top, left = index, 0
        rows.append((top, left, text, score))
    rows.sort(key=lambda row: (round(row[0] / 12), row[1]))
    return "\n".join(row[2] for row in rows), average([row[3] for row in rows])


def image_to_text(ocr, image):
    import numpy as np

    result = ocr.predict(np.asarray(image.convert("RGB")))
    for item in result:
        text, confidence = text_from_ocr_result(item)
        if text:
            return text, confidence
    return "", 0.0


def emit_page(source, page_index, text, confidence, error=None):
    value = {
        "kind": "page",
        "fileId": source["id"],
        "name": source["name"],
        "pageIndex": page_index,
        "text": text,
        "confidence": confidence,
    }
    if error:
        value["error"] = error
    emit(value)


def recognize_source(source, ocr):
    path = Path(source["path"])
    extension = path.suffix.lower()

    if extension in {".txt", ".text", ".md"}:
        text = read_text_file(path)
        emit_page(source, 0, text, 1.0 if text else 0.0, None if text else "文本文件为空或无法读取。")
        return 1

    if extension == ".docx":
        text = extract_docx(path)
        emit_page(source, 0, text, 1.0 if text else 0.0, None if text else "文档中没有提取到文字。")
        return 1

    if extension == ".doc":
        result = subprocess.run(["antiword", str(path)], capture_output=True, text=True, check=True)
        text = result.stdout.strip()
        emit_page(source, 0, text, 1.0 if text else 0.0, None if text else "文档中没有提取到文字。")
        return 1

    if extension == ".pdf":
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(path))
        for index, page in enumerate(pdf):
            try:
                text_page = page.get_textpage()
                text = text_page.get_text_range().strip()
                text_page.close()
                confidence = 1.0 if text else 0.0
                if not text:
                    bitmap = page.render(scale=2.2)
                    image = bitmap.to_pil()
                    text, confidence = image_to_text(ocr, image)
                    bitmap.close()
                emit_page(source, index, text, confidence, None if text else "本页没有识别到可读文字，请检查原图。")
            except Exception as error:  # Keep later pages flowing after a damaged page.
                emit_page(source, index, "", 0.0, f"PDF 第 {index + 1} 页处理失败：{error}")
        return len(pdf)

    if extension in {".heic", ".heif"}:
        from pillow_heif import register_heif_opener

        register_heif_opener()

    from PIL import Image, ImageOps

    with Image.open(path) as original:
        image = ImageOps.exif_transpose(original)
        text, confidence = image_to_text(ocr, image)
    emit_page(source, 0, text, confidence, None if text else "图片中没有识别到可读文字，请检查原图。")
    return 1


def main():
    try:
        sources = json.load(sys.stdin)
        if not isinstance(sources, list):
            raise ValueError("处理清单必须是数组。")
    except Exception as error:
        emit({"kind": "fatal", "error": f"无法读取处理清单：{error}"})
        return 2

    try:
        from paddleocr import PaddleOCR

        ocr = PaddleOCR(
            lang=os.environ.get("SCRIBE_OCR_LANG", "en"),
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
        )
    except Exception as error:
        emit({"kind": "fatal", "error": f"OCR 引擎初始化失败：{error}"})
        return 2

    count = 0
    for source in sources:
        try:
            count += recognize_source(source, ocr)
        except Exception as error:
            emit_page(source, 0, "", 0.0, f"{source.get('name', '文件')}处理失败：{error}")
            count += 1
    emit({"kind": "done", "count": count})
    return 0


if __name__ == "__main__":
    sys.exit(main())
